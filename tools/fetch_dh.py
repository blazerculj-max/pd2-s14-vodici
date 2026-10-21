"""DarkHumilityjeva S14 tier lista (Google Sheet) → data/dh/tierlist.json.

Tabela nima API-ja, zato jo prenesemo kot xlsx in razberemo bloke po glavah:
tier lista po razredih (B–G, H–M), seznami top 10 (N–R), priporočila za starterje in bosse
(S–AB, po dva stolpca) ter mape z imunitetami (AC–AI)."""
import datetime
import json
import os
import re
import subprocess
from concurrent.futures import ThreadPoolExecutor

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "dh")
SHEET_ID = "1DLkTz-3gtlBgoyxxAot60Ro_76ct3Oz8Bmyhu7Hbwo4"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/htmlview"
CLASSES = {"Sorceress": "sorceress", "Druid": "druid", "Assassin": "assassin", "Barbarian": "barbarian",
           "Amazon": "amazon", "Necromancer": "necromancer", "Paladin": "paladin"}
REC_KEYS = {"Starter": "starter", "Uber Tristram": "uber", "Dclone": "dclone", "Rathma": "rathma", "Lucion": "lucion"}
MAP_EL = {"Fire": "fire", "Cold": "cold", "Lightning": "light", "Poison": "poison", "Physical": "phys", "Magic": "magic"}


def download(path):
    os.makedirs(OUT, exist_ok=True)
    subprocess.run(["curl", "-s", "-L", "-f", "-A", "Mozilla/5.0", "-o", path,
                    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"], check=True)


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else None


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def parse(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    grid, links = {}, {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                grid[(c.column_letter, c.row)] = text(c.value)
            if c.hyperlink and c.hyperlink.target:
                links[(c.column_letter, c.row)] = c.hyperlink.target
    maxr = ws.max_row
    g = lambda col, r: grid.get((col, r))

    # tier lista: dva bloka razredov (ime, tier, fortified, budget, HC, T1–T2)
    tiers = []
    for cols in (("B", "C", "D", "E", "F", "G"), ("H", "I", "J", "K", "L", "M")):
        cls, tree = None, None
        for r in range(1, maxr + 1):
            name, tier = g(cols[0], r), g(cols[1], r)
            if not name:
                continue
            if tier == "Tier" and name.endswith(" Builds"):
                cls, tree = CLASSES.get(name[:-7]), None
            elif cls and not tier:
                tree = name
            elif cls and tier:
                tiers.append({"cls": cls, "tree": tree, "name": name, "tier": tier, "fort": g(cols[2], r),
                              "budget": g(cols[3], r), "hc": g(cols[4], r), "t12": g(cols[5], r),
                              "video": links.get((cols[0], r))})

    # seznami v stolpcih N–R: glava brez »#«, nato vrstice #1, #2 …
    lists, cur = [], None
    for r in range(1, maxr + 1):
        n = g("N", r)
        if not n:
            continue
        if not n.startswith("#"):
            cur = {"key": slug(n), "title": n, "items": []}
            lists.append(cur)
        elif cur is not None:
            cur["items"].append({"rank": n, "name": g("O", r), "from": g("P", r), "to": g("R", r) if g("Q", r) else None,
                                 "tier": None if g("Q", r) else g("R", r)})

    # priporočila: pari stolpcev, glava v 1. vrstici, nato razred ali »Overall …«, nato #1, #2 …
    recs = {}
    for a, b in (("S", "T"), ("U", "V"), ("W", "X"), ("Y", "Z"), ("AA", "AB")):
        head = g(a, 1) or ""
        key = next((v for k, v in REC_KEYS.items() if k.lower() in head.lower()), slug(head))
        block, grp = {"title": head, "classes": {}, "overall": []}, None
        for r in range(2, maxr + 1):
            x = g(a, r)
            if not x:
                continue
            if x.startswith("#"):
                if grp is not None and g(b, r):
                    grp.append(g(b, r))
            elif x in CLASSES:
                grp = block["classes"].setdefault(CLASSES[x], [])
            elif x.lower().startswith("overall"):
                block["overall_title"] = x
                grp = block["overall"]
        recs[key] = block

    # mape: »Tier N« je glava, sledijo mape z imunitetami (številke) ali opombo (unique mape)
    header = {g(c, r): c for r in range(1, 10) for c in ("AD", "AE", "AF", "AG", "AH", "AI") if g(c, r) in MAP_EL}
    maps, tier = [], None
    for r in range(1, maxr + 1):
        name = g("AC", r)
        if not name or name in ("Maps", "S14 Map Tiers/Immunities"):
            continue
        if re.match(r"^Tier \d", name):
            tier = name
            continue
        if not tier:
            continue
        imm = {}
        for el, c in header.items():
            v = g(c, r)
            if v and re.fullmatch(r"\d+(\.0)?", v):
                imm[MAP_EL[el]] = int(float(v))
        note = g("AD", r) if not imm and g("AD", r) and not re.fullmatch(r"\d+(\.0)?", g("AD", r)) else None
        maps.append({"tier": tier, "name": re.sub(r"\s*\((Uniq|New)\)", "", name).strip(), "unique": "(Uniq)" in name,
                     "new": "(New)" in name, "imm": imm, "note": note})

    # povezave do drugih DarkHumilityjevih vodičev (stolpec A)
    extra = [{"label": g("A", r), "url": links[("A", r)]} for r in range(2, maxr + 1)
             if ("A", r) in links and g("A", r) and "twitch" not in links[("A", r)]]

    return {"url": SHEET_URL, "title": g("B", 1), "fetched": datetime.date.today().isoformat(),
            "tiers": tiers, "lists": lists, "recs": recs, "maps": maps, "links": extra}


def video_meta(url):
    """Naslov, datum in sezona YouTube videa (iz strani videa; API ni potreben)."""
    m = re.search(r"v=([\w-]{11})", url or "")
    if not m:
        return None
    r = subprocess.run(["curl", "-s", "-L", "-A", "Mozilla/5.0", "-H", "Accept-Language: en-US,en;q=0.9", "-b", "CONSENT=YES+1",
                        f"https://www.youtube.com/watch?v={m.group(1)}"], capture_output=True, text=True, errors="ignore")
    t = re.search(r'"title":"((?:[^"\\]|\\.)*)","lengthSeconds"', r.stdout)
    d = re.search(r'"publishDate":"([^"]+)"', r.stdout)
    title = json.loads('"' + t.group(1) + '"') if t else None
    season = re.search(r"\bS(?:eason\s*)?(\d{1,2})\b", title or "")
    return {"id": m.group(1), "title": title, "date": d.group(1)[:10] if d else None,
            "season": int(season.group(1)) if season else None}


def main():
    path = os.path.join(OUT, "sheet.xlsx")
    download(path)
    data = parse(path)
    urls = [t["video"] for t in data["tiers"] if t["video"]] + [l["url"] for l in data["links"] if "youtube" in l["url"]]
    with ThreadPoolExecutor(8) as ex:
        data["videos"] = {v["id"]: v for v in ex.map(video_meta, urls) if v}
    json.dump(data, open(os.path.join(OUT, "tierlist.json"), "w"), indent=1, ensure_ascii=False)
    print(f'DarkHumility: {len(data["tiers"])} buildov, {sum(1 for t in data["tiers"] if t["video"])} videov, '
          f'{len(data["lists"])} seznamov, {len(data["recs"])} priporočil, {len(data["maps"])} map')


if __name__ == "__main__":
    main()
